from rest_framework.decorators import api_view
from rest_framework.response import Response
from django.http import HttpResponse
from django.db import models
from django.db.models import Count, Q, F, Value, Subquery, OuterRef, IntegerField
from django.db.models.functions import ExtractYear, Coalesce, Lower
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import quote_sheetname
from openpyxl.worksheet.hyperlink import Hyperlink
from django.utils import timezone
from .models import (
    Researcher,
    Article,
    ArticleCategory,
    Thesis,
    Book,
    BookChapter,
    Country,
    ResearchGroup,
    ScientificEvent,
    ResearchLine,
    ArticleAuthor,
    BookAuthor,
    ChapterAuthor,
    ThesisTutor,
    ThesisStudent,
    EventInstitution,
    GenericProduct,
    GenericProductAuthor,
    GroupMember,
)
from .infrastructure.utils.country_coords import COUNTRY_COORDINATES
import re
import unicodedata

def remove_accents(input_str):
    if not input_str: return ""
    nfkd_form = unicodedata.normalize('NFKD', input_str)
    return "".join([c for c in nfkd_form if not unicodedata.category(c).startswith('Mn')])


def normalize_identifier(value):
    if not value:
        return ""
    return re.sub(r'[^a-z0-9]', '', remove_accents(value).lower())


def normalize_country_key(value):
    if not value:
        return ""
    normalized = remove_accents(value).lower()
    normalized = re.sub(r'[^a-z0-9]+', ' ', normalized)
    return re.sub(r'\s+', ' ', normalized).strip()


def normalize_doi(value):
    if not value:
        return ""
    return normalize_identifier(value)


COUNTRY_ALIAS_GROUPS = {
    'estados unidos': {
        'estados unidos',
        'estados unidos de america',
        'usa',
        'u s a',
        'eeuu',
        'ee uu',
        'ee.uu',
        'united states',
        'us',
    },
    'reino unido': {
        'reino unido',
        'inglaterra',
        'united kingdom',
        'uk',
        'great britain',
        'britain',
        'gran bretana',
    },
}

COUNTRY_CANONICAL_DISPLAY_NAMES = {
    'estados unidos': 'Estados Unidos',
    'reino unido': 'Reino Unido',
}

COUNTRY_ALIAS_LOOKUP = {
    normalize_country_key(alias): canonical_key
    for canonical_key, aliases in COUNTRY_ALIAS_GROUPS.items()
    for alias in aliases
}


def canonicalize_country_name(value):
    normalized_key = normalize_country_key(value)
    return COUNTRY_ALIAS_LOOKUP.get(normalized_key, normalized_key)


def get_country_display_name(value):
    canonical_key = canonicalize_country_name(value)
    return COUNTRY_CANONICAL_DISPLAY_NAMES.get(canonical_key, (value or '').title())


def resolve_country_group(country_id):
    raw_value = (country_id or '').strip()
    all_countries = list(Country.objects.all())

    try:
        import uuid
        uuid_obj = uuid.UUID(raw_value)
        country = next((item for item in all_countries if item.id == uuid_obj), None)
        if country:
            canonical_key = canonicalize_country_name(country.name)
            matched = [item for item in all_countries if canonicalize_country_name(item.name) == canonical_key]
            return matched, COUNTRY_CANONICAL_DISPLAY_NAMES.get(canonical_key, country.name.title()), canonical_key
    except (ValueError, TypeError):
        pass

    normalized_input = canonicalize_country_name(raw_value.replace('-', ' '))
    if not normalized_input:
        return [], None, None

    matched = [item for item in all_countries if canonicalize_country_name(item.name) == normalized_input]
    if matched:
        return matched, COUNTRY_CANONICAL_DISPLAY_NAMES.get(normalized_input, matched[0].name.title()), normalized_input

    if len(normalized_input) > 3:
        for country in all_countries:
            country_key = canonicalize_country_name(country.name)
            if normalized_input in country_key or country_key in normalized_input:
                matched = [item for item in all_countries if canonicalize_country_name(item.name) == country_key]
                return matched, COUNTRY_CANONICAL_DISPLAY_NAMES.get(country_key, country.name.title()), country_key

    return [], None, normalized_input


ACCENTED_CHARS = 'áàäâãéèëêíìïîóòöôõúùüûñç'
PLAIN_CHARS = 'aaaaaeeeeiiiiooooouuuunc'


def normalized_text_annotation(field_name):
    """Builds a lowercase accent-insensitive SQL expression for text search."""
    return Lower(
        models.Func(
            Lower(Coalesce(F(field_name), Value(''), output_field=models.TextField())),
            Value(ACCENTED_CHARS),
            Value(PLAIN_CHARS),
            function='TRANSLATE',
            output_field=models.TextField(),
        )
    )


def unique_by_id(items):
    """Deduplicates response rows preserving original order."""
    seen = set()
    unique_items = []
    for item in items:
        item_id = item.get('id')
        if item_id in seen:
            continue
        seen.add(item_id)
        unique_items.append(item)
    return unique_items


def serialize_optional_datetime(value):
    if not value:
        return None
    localized = timezone.localtime(value) if timezone.is_aware(value) else value
    return localized.date().isoformat()


GENERIC_EXTRA_DATA_PREFERRED_KEYS = (
    'tipo',
    'ciudad',
    'pais',
    'ambito',
    'disponibilidad',
    'idioma',
    'institucionSolicitante',
    'institucionServicio',
    'institucionFinanciadora',
    'fechaEnvio',
    'mes',
    'numeroConsecutivoConcepto',
    'numeroContrato',
)


def humanize_extra_data_key(value):
    labels = {
        'tipo': 'Tipo',
        'ciudad': 'Ciudad',
        'pais': 'País',
        'ambito': 'Ámbito',
        'disponibilidad': 'Disponibilidad',
        'idioma': 'Idioma',
        'fechaEnvio': 'Fecha de envío',
        'mes': 'Mes',
        'numeroContrato': 'Número de contrato',
        'numeroConsecutivoConcepto': 'Número consecutivo',
        'institucionSolicitante': 'Institución solicitante',
        'institucionServicio': 'Institución de servicio',
        'institucionFinanciadora': 'Institución financiadora',
    }
    if value in labels:
        return labels[value]
    return re.sub(r'(?<!^)([A-Z])', r' \1', str(value or '').replace('_', ' ')).strip().capitalize()


def stringify_extra_data_value(value):
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'Sí' if value else 'No'
    if isinstance(value, (list, tuple, set)):
        return ', '.join([stringify_extra_data_value(item) for item in value if item not in (None, '')])
    if isinstance(value, dict):
        return '; '.join(
            f"{humanize_extra_data_key(key)}: {stringify_extra_data_value(item)}"
            for key, item in value.items()
            if item not in (None, '')
        )
    return str(value).strip().rstrip(':;,')


def ordered_generic_extra_data_keys(rows):
    keys = set()
    for row in rows:
        extra_data = row.get('extraData') or {}
        if not isinstance(extra_data, dict):
            continue
        for key, value in extra_data.items():
            if key == 'raw_text' or value in (None, ''):
                continue
            keys.add(key)

    def sort_key(key):
        try:
            return (0, GENERIC_EXTRA_DATA_PREFERRED_KEYS.index(key))
        except ValueError:
            return (1, humanize_extra_data_key(key).lower())

    return sorted(keys, key=sort_key)


def parse_optional_year_param(value):
    if value in (None, ''):
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def get_event_year_from_payload(item):
    source = item.get('startDate') or item.get('endDate')
    if not source:
        return None
    try:
        return int(str(source)[:4])
    except (TypeError, ValueError):
        return None


def within_year_range(year, from_year=None, to_year=None):
    if not from_year and not to_year:
        return True
    if not year:
        return False
    if from_year and year < from_year:
        return False
    if to_year and year > to_year:
        return False
    return True


def filter_products_payload_by_year_range(products, from_year=None, to_year=None):
    if not from_year and not to_year:
        return products

    return {
        **products,
        'articles': [
            item for item in products.get('articles', [])
            if within_year_range(item.get('year'), from_year, to_year)
        ],
        'books': [
            item for item in products.get('books', [])
            if within_year_range(item.get('year'), from_year, to_year)
        ],
        'bookChapters': [
            item for item in products.get('bookChapters', [])
            if within_year_range(item.get('year'), from_year, to_year)
        ],
        'theses': [
            item for item in products.get('theses', [])
            if within_year_range(item.get('year'), from_year, to_year)
        ],
        'events': [
            item for item in products.get('events', [])
            if within_year_range(get_event_year_from_payload(item), from_year, to_year)
        ],
        'generic': [
            item for item in products.get('generic', [])
            if within_year_range(item.get('year'), from_year, to_year)
        ],
    }


def get_available_years_from_products(products):
    years = set()
    current_year = timezone.localtime().year

    def register_year(value):
        if not value:
            return
        if 1900 <= value <= current_year + 1:
            years.add(value)

    for item in products.get('articles', []):
        register_year(item.get('year'))
    for item in products.get('books', []):
        register_year(item.get('year'))
    for item in products.get('bookChapters', []):
        register_year(item.get('year'))
    for item in products.get('theses', []):
        register_year(item.get('year'))
    for item in products.get('events', []):
        register_year(get_event_year_from_payload(item))
    for item in products.get('generic', []):
        register_year(item.get('year'))
    return sorted(years, reverse=True)


def filter_group_detail_payload_by_year_range(data, from_year=None, to_year=None):
    if not from_year and not to_year:
        return data

    return {
        **data,
        'products': filter_products_payload_by_year_range(data['products'], from_year, to_year),
    }


def describe_export_year_filter(from_year=None, to_year=None):
    if not from_year and not to_year:
        return 'Todos los años'
    if from_year and to_year:
        return f'Desde {from_year} hasta {to_year}'
    if from_year:
        return f'Desde {from_year} hasta el presente'
    return f'Desde el año más antiguo hasta {to_year}'


def _normalize_text_key(value):
    return normalize_identifier(str(value or ''))


def _append_unique_thesis_row(target_rows, row, related_field):
    key = (
        _normalize_text_key(row.get('title')),
        row.get('year') or None,
        _normalize_text_key(row.get('thesisType')),
        _normalize_text_key(row.get('institution')),
    )

    for existing in target_rows:
        existing_key = (
            _normalize_text_key(existing.get('title')),
            existing.get('year') or None,
            _normalize_text_key(existing.get('thesisType')),
            _normalize_text_key(existing.get('institution')),
        )
        if existing_key == key:
            merged = list(dict.fromkeys((existing.get(related_field) or []) + (row.get(related_field) or [])))
            existing[related_field] = merged
            return

    row[related_field] = list(dict.fromkeys(row.get(related_field) or []))
    target_rows.append(row)


def identifier_annotation(field_name):
    # Normalizes DOI/ISSN/ISBN in DB (remove non-alphanumeric chars).
    return Lower(
        models.Func(
            Coalesce(F(field_name), Value(''), output_field=models.TextField()),
            Value('[^a-zA-Z0-9]'),
            Value(''),
            Value('g'),
            function='REGEXP_REPLACE',
            output_field=models.TextField(),
        )
    )


@api_view(['GET'])
def global_search(request):
    query = (request.query_params.get('q') or "").strip()
    if not query:
        return Response(
            {"error": "Query parameter 'q' is required."},
            status=400,
        )

    per_type_limit = request.query_params.get('limit', '10')
    try:
        per_type_limit = max(1, min(int(per_type_limit), 50))
    except ValueError:
        per_type_limit = 10

    query_identifier = normalize_identifier(query)
    query_doi = normalize_doi(query)
    query_normalized = remove_accents(query).lower()

    article_ids_by_author = ArticleAuthor.objects.annotate(
        normalized_author_name=normalized_text_annotation('author_name'),
        normalized_researcher_name=normalized_text_annotation('researcher__name'),
    ).filter(
        Q(normalized_author_name__icontains=query_normalized) |
        Q(normalized_researcher_name__icontains=query_normalized)
    ).values_list('article_id', flat=True)

    book_ids_by_author = BookAuthor.objects.annotate(
        normalized_author_name=normalized_text_annotation('author_name'),
        normalized_researcher_name=normalized_text_annotation('researcher__name'),
    ).filter(
        Q(normalized_author_name__icontains=query_normalized) |
        Q(normalized_researcher_name__icontains=query_normalized)
    ).values_list('book_id', flat=True)

    chapter_ids_by_author = ChapterAuthor.objects.annotate(
        normalized_author_name=normalized_text_annotation('author_name'),
        normalized_researcher_name=normalized_text_annotation('researcher__name'),
    ).filter(
        Q(normalized_author_name__icontains=query_normalized) |
        Q(normalized_researcher_name__icontains=query_normalized)
    ).values_list('chapter_id', flat=True)

    researchers_qs = (
        Researcher.objects.annotate(
            normalized_name=normalized_text_annotation('name'),
        )
        .filter(
            Q(normalized_name__icontains=query_normalized) |
            Q(name__icontains=query) |
            Q(code_rh__icontains=query)
        )
        .order_by('name')[:per_type_limit]
    )

    groups_qs = (
        ResearchGroup.objects.annotate(
            normalized_name=normalized_text_annotation('name'),
        )
        .filter(
            Q(normalized_name__icontains=query_normalized) |
            Q(name__icontains=query) |
            Q(code__icontains=query) |
            Q(leader__icontains=query)
        )
        .order_by('name')[:per_type_limit]
    )

    articles_qs = (
        Article.objects.select_related('journal', 'group')
        .annotate(
            normalized_issn=identifier_annotation('issn'),
            normalized_journal_issn=identifier_annotation('journal__issn'),
            normalized_doi=identifier_annotation('doi'),
            normalized_title=normalized_text_annotation('title'),
            normalized_journal_name=normalized_text_annotation('journal__name'),
        )
        .filter(
            Q(normalized_title__icontains=query_normalized) |
            Q(normalized_journal_name__icontains=query_normalized) |
            Q(id__in=article_ids_by_author) |
            Q(title__icontains=query) |
            Q(doi__icontains=query) |
            Q(issn__icontains=query) |
            Q(journal__name__icontains=query) |
            Q(journal__issn__icontains=query) |
            Q(normalized_issn=query_identifier) |
            Q(normalized_journal_issn=query_identifier) |
            Q(normalized_doi=query_doi)
        )
        .distinct()
        .order_by('-year', 'title')[:per_type_limit]
    )

    books_qs = (
        Book.objects.select_related('publisher', 'group')
        .annotate(
            normalized_isbn=identifier_annotation('isbn'),
            normalized_title=normalized_text_annotation('title'),
        )
        .filter(
            Q(normalized_title__icontains=query_normalized) |
            Q(id__in=book_ids_by_author) |
            Q(title__icontains=query) |
            Q(isbn__icontains=query) |
            Q(normalized_isbn=query_identifier)
        )
        .distinct()
        .order_by('-year', 'title')[:per_type_limit]
    )

    chapters_qs = (
        BookChapter.objects.select_related('publisher', 'group')
        .annotate(
            normalized_isbn=identifier_annotation('isbn'),
            normalized_chapter_title=normalized_text_annotation('chapter_title'),
            normalized_book_title=normalized_text_annotation('book_title'),
        )
        .filter(
            Q(normalized_chapter_title__icontains=query_normalized) |
            Q(normalized_book_title__icontains=query_normalized) |
            Q(id__in=chapter_ids_by_author) |
            Q(chapter_title__icontains=query) |
            Q(book_title__icontains=query) |
            Q(isbn__icontains=query) |
            Q(normalized_isbn=query_identifier)
        )
        .distinct()
        .order_by('-year', 'chapter_title')[:per_type_limit]
    )

    researchers = [
        {
            'id': str(r.id),
            'name': r.name,
            'category': r.category,
            'codeRh': r.code_rh,
            'detailPath': f'/investigadores/{r.id}',
        }
        for r in researchers_qs
    ]

    groups = [
        {
            'id': str(g.id),
            'name': g.name,
            'code': g.code,
            'leader': g.leader,
            'category': g.category,
            'detailPath': f'/groups/{g.id}',
        }
        for g in groups_qs
    ]

    articles = [
        {
            'id': str(a.id),
            'title': a.title,
            'year': a.year,
            'doi': a.doi,
            'issn': a.issn or (a.journal.issn if a.journal else None),
            'journal': a.journal.name if a.journal else None,
            'groupName': a.group.name if a.group else None,
            'detailPath': f'/articulos/{a.id}',
        }
        for a in articles_qs
    ]

    books = [
        {
            'id': str(b.id),
            'title': b.title,
            'year': b.year,
            'isbn': b.isbn,
            'publisher': b.publisher.name if b.publisher else None,
            'groupName': b.group.name if b.group else None,
            'detailPath': f'/libros/{b.id}',
        }
        for b in books_qs
    ]

    book_chapters = [
        {
            'id': str(c.id),
            'chapterTitle': c.chapter_title,
            'bookTitle': c.book_title,
            'year': c.year,
            'isbn': c.isbn,
            'publisher': c.publisher.name if c.publisher else None,
            'groupName': c.group.name if c.group else None,
            'detailPath': f'/capitulos/{c.id}',
        }
        for c in chapters_qs
    ]

    # Defensive dedupe for joins with authors and other related records.
    researchers = unique_by_id(researchers)
    groups = unique_by_id(groups)
    articles = unique_by_id(articles)
    books = unique_by_id(books)
    book_chapters = unique_by_id(book_chapters)

    return Response(
        {
            'query': query,
            'totals': {
                'researchers': len(researchers),
                'groups': len(groups),
                'articles': len(articles),
                'books': len(books),
                'bookChapters': len(book_chapters),
                'all': len(researchers) + len(groups) + len(articles) + len(books) + len(book_chapters),
            },
            'results': {
                'researchers': researchers,
                'groups': groups,
                'articles': articles,
                'books': books,
                'bookChapters': book_chapters,
            },
        }
    )


@api_view(['GET'])
def get_researcher_detail(request, researcher_id):
    researcher = Researcher.objects.filter(id=researcher_id).first()
    if not researcher:
        return Response({'error': 'Researcher not found'}, status=404)

    return Response(_build_researcher_detail_payload(researcher))


def _build_researcher_detail_payload(researcher):
    groups = ResearchGroup.objects.filter(members__researcher=researcher).distinct().order_by('name')

    articles = Article.objects.filter(authors__researcher=researcher).select_related('journal', 'group', 'product_type_obj').distinct().order_by('-year', 'title')
    books = Book.objects.filter(authors__researcher=researcher).select_related('publisher', 'group').distinct().order_by('-year', 'title')
    chapters = BookChapter.objects.filter(authors__researcher=researcher).select_related('publisher', 'group').distinct().order_by('-year', 'chapter_title')

    article_ids = [a.id for a in articles]
    book_ids = [b.id for b in books]
    chapter_ids = [c.id for c in chapters]

    article_author_counts = dict(
        ArticleAuthor.objects.filter(article_id__in=article_ids)
        .values('article_id')
        .annotate(total=Count('id'))
        .values_list('article_id', 'total')
    )
    book_author_counts = dict(
        BookAuthor.objects.filter(book_id__in=book_ids)
        .values('book_id')
        .annotate(total=Count('id'))
        .values_list('book_id', 'total')
    )
    chapter_author_counts = dict(
        ChapterAuthor.objects.filter(chapter_id__in=chapter_ids)
        .values('chapter_id')
        .annotate(total=Count('id'))
        .values_list('chapter_id', 'total')
    )

    article_categories_by_id = {}
    article_categories_qs = (
        ArticleCategory.objects.filter(article_id__in=article_ids, source__name__iexact='PUBLINDEX')
        .select_related('source')
        .values('article_id', 'year', 'category', 'source__name')
    )
    for row in article_categories_qs:
        article_categories_by_id.setdefault(row['article_id'], []).append(row)

    def resolve_article_category(article):
        rows = article_categories_by_id.get(article.id, [])
        if not rows:
            return {
                'category': 'N/R',
                'source': None,
            }

        target_year = article.year
        if target_year is not None:
            filtered = [r for r in rows if r.get('year') == target_year]
            if filtered:
                rows = filtered

        publindex_rows = [row for row in rows if row.get('category')]
        if publindex_rows:
            selected = max(publindex_rows, key=lambda row: row.get('year') or 0)
            return {
                'category': str(selected.get('category')),
                'source': 'Publindex',
            }

        return {
            'category': 'N/R',
            'source': None,
        }

    thesis_tutor_links = ThesisTutor.objects.filter(researcher=researcher).select_related('thesis').order_by('-thesis__year', 'thesis__title', 'order')
    thesis_student_links = ThesisStudent.objects.filter(researcher=researcher).select_related('thesis').order_by('-thesis__year', 'thesis__title')

    tutor_thesis_ids = [link.thesis_id for link in thesis_tutor_links]
    student_thesis_ids = [link.thesis_id for link in thesis_student_links]
    thesis_ids = list(set(tutor_thesis_ids + student_thesis_ids))

    thesis_map = {
        thesis.id: thesis
        for thesis in Thesis.objects.filter(id__in=thesis_ids).select_related('institution_obj').order_by('-year', 'title')
    }

    thesis_students_by_thesis = {}
    thesis_students_qs = ThesisStudent.objects.filter(thesis_id__in=thesis_ids).order_by('order', 'student_name')
    for student in thesis_students_qs:
        thesis_students_by_thesis.setdefault(student.thesis_id, []).append(student.student_name)

    thesis_as_director = []
    thesis_as_cotutor = []
    thesis_as_student = []

    for link in thesis_tutor_links:
        thesis = thesis_map.get(link.thesis_id)
        if not thesis:
            continue
        thesis_row = {
            'id': str(thesis.id),
            'title': thesis.title,
            'year': thesis.year,
            'thesisType': thesis.thesis_type,
            'institution': thesis.institution_obj.name if thesis.institution_obj else thesis.institution,
            'students': thesis_students_by_thesis.get(thesis.id, []),
        }
        if link.order == 0:
            _append_unique_thesis_row(thesis_as_director, thesis_row, 'students')
        else:
            _append_unique_thesis_row(thesis_as_cotutor, thesis_row, 'students')

    for link in thesis_student_links:
        thesis = thesis_map.get(link.thesis_id)
        if not thesis:
            continue
        _append_unique_thesis_row(
            thesis_as_student,
            {
                'id': str(thesis.id),
                'title': thesis.title,
                'year': thesis.year,
                'thesisType': thesis.thesis_type,
                'institution': thesis.institution_obj.name if thesis.institution_obj else thesis.institution,
                'tutors': [
                    t.tutor_name
                    for t in ThesisTutor.objects.filter(thesis=thesis).order_by('order', 'tutor_name')
                ],
            },
            'tutors',
        )

    def unique_articles(items):
        seen = set()
        unique_rows = []
        for item in items:
            key = (
                _normalize_text_key(item.title),
                item.year or None,
                _normalize_text_key(item.doi),
                _normalize_text_key(item.issn or (item.journal.issn if item.journal else None)),
            )
            if key in seen:
                continue
            seen.add(key)
            unique_rows.append(item)
        return unique_rows

    def unique_books(items):
        seen = set()
        unique_rows = []
        for item in items:
            key = (
                _normalize_text_key(item.title),
                item.year or None,
                _normalize_text_key(item.isbn),
            )
            if key in seen:
                continue
            seen.add(key)
            unique_rows.append(item)
        return unique_rows

    def unique_chapters(items):
        seen = set()
        unique_rows = []
        for item in items:
            key = (
                _normalize_text_key(item.chapter_title),
                _normalize_text_key(item.book_title),
                item.year or None,
                _normalize_text_key(item.isbn),
            )
            if key in seen:
                continue
            seen.add(key)
            unique_rows.append(item)
        return unique_rows

    unique_article_rows = unique_articles(list(articles))[:200]
    unique_book_rows = unique_books(list(books))[:200]
    unique_chapter_rows = unique_chapters(list(chapters))[:200]

    return {
        'id': str(researcher.id),
        'name': researcher.name,
        'codeRh': researcher.code_rh,
        'category': researcher.category,
        'educationLevel': researcher.education_level,
        'city': researcher.city,
        'department': researcher.department,
        'university': researcher.university,
        'cvlacUrl': researcher.cvlac_url,
        'groups': [
            {
                'id': str(group.id),
                'name': group.name,
                'code': group.code,
            }
            for group in groups
        ],
        'products': {
            'articles': [
                {
                    'id': str(article.id),
                    'title': article.title,
                    'year': article.year,
                    'category': resolved_category['category'],
                    'categorySource': resolved_category['source'],
                    'issn': article.issn or (article.journal.issn if article.journal else None),
                    'doi': article.doi,
                    'researchersCount': article_author_counts.get(article.id, 0),
                    'detailPath': f'/articulos/{article.id}',
                }
                for article in unique_article_rows
                for resolved_category in [resolve_article_category(article)]
            ],
            'books': [
                {
                    'id': str(book.id),
                    'title': book.title,
                    'year': book.year,
                    'isbn': book.isbn,
                    'publisher': book.publisher.name if book.publisher else None,
                    'researchersCount': book_author_counts.get(book.id, 0),
                    'detailPath': f'/libros/{book.id}',
                }
                for book in unique_book_rows
            ],
            'bookChapters': [
                {
                    'id': str(chapter.id),
                    'chapterTitle': chapter.chapter_title,
                    'bookTitle': chapter.book_title,
                    'year': chapter.year,
                    'isbn': chapter.isbn,
                    'publisher': chapter.publisher.name if chapter.publisher else None,
                    'researchersCount': chapter_author_counts.get(chapter.id, 0),
                    'detailPath': f'/capitulos/{chapter.id}',
                }
                for chapter in unique_chapter_rows
            ],
        },
        'theses': {
            'asDirector': thesis_as_director,
            'asCotutor': thesis_as_cotutor,
            'asStudent': thesis_as_student,
        },
    }


@api_view(['GET'])
def export_researcher_detail_excel(request, researcher_id):
    researcher = Researcher.objects.filter(id=researcher_id).first()
    if not researcher:
        return Response({'error': 'Researcher not found'}, status=404)

    data = _build_researcher_detail_payload(researcher)

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    # Frontend primary tone: hsl(162 47% 18%) ~= #184336
    header_fill = PatternFill(start_color='184336', end_color='184336', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    title_font = Font(size=18, bold=True, color='184336')
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB'),
    )

    def style_table_header(sheet, total_columns):
        for col_idx in range(1, total_columns + 1):
            cell = sheet.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = thin_border
        sheet.freeze_panes = 'A2'

    def auto_fit_columns(sheet, min_width=12, max_width=60):
        for col_cells in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col_cells[0].column)
            for cell in col_cells:
                cell_value = '' if cell.value is None else str(cell.value)
                max_len = max(max_len, len(cell_value))
            sheet.column_dimensions[col_letter].width = max(min(max_len + 2, max_width), min_width)

    cover_sheet = workbook.create_sheet('Portada', 0)
    code_rh = (researcher.code_rh or 'SIN-CODIGO').strip()
    researcher_name = (researcher.name or 'Investigador').strip()

    cover_sheet.merge_cells('A1:E1')
    cover_sheet['A1'] = 'Ficha de Investigador'
    cover_sheet['A1'].font = title_font
    cover_sheet['A1'].alignment = Alignment(horizontal='center', vertical='center')

    cover_sheet['A3'] = 'Código RH'
    cover_sheet['B3'] = code_rh
    cover_sheet['A4'] = 'Nombre del Investigador'
    cover_sheet['B4'] = researcher_name
    cover_sheet['A5'] = 'Categoría'
    cover_sheet['B5'] = researcher.category or 'N/A'
    cover_sheet['A6'] = 'Formación'
    cover_sheet['B6'] = researcher.education_level or 'N/A'
    cover_sheet['A7'] = 'Fecha de generación'
    cover_sheet['B7'] = timezone.localtime().strftime('%Y-%m-%d %H:%M')

    for row in range(3, 8):
        label_cell = cover_sheet[f'A{row}']
        value_cell = cover_sheet[f'B{row}']
        label_cell.font = Font(bold=True, color='184336')
        label_cell.fill = PatternFill(start_color='E6F1EE', end_color='E6F1EE', fill_type='solid')
        label_cell.alignment = Alignment(horizontal='left', vertical='center')
        label_cell.border = thin_border
        value_cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        value_cell.border = thin_border

    cover_sheet.column_dimensions['A'].width = 28
    cover_sheet.column_dimensions['B'].width = 60
    cover_sheet.row_dimensions[1].height = 30

    ws_articles = workbook.create_sheet('Articulos')
    ws_articles.append(['Título', 'Año', 'Categoría', 'Fuente', 'ISSN', 'DOI', 'No. Investigadores'])
    for row in data['products']['articles']:
        category_value = (row.get('category') or '').strip() if isinstance(row.get('category'), str) else row.get('category')
        category_value = category_value if category_value else 'N/R'
        source_value = (row.get('categorySource') or '').strip() if isinstance(row.get('categorySource'), str) else row.get('categorySource')
        source_value = source_value if source_value else 'N/R'

        if category_value == 'N/R':
            source_value = 'N/R'

        ws_articles.append([
            row['title'],
            row.get('year'),
            category_value,
            source_value,
            row.get('issn'),
            row.get('doi'),
            row.get('researchersCount'),
        ])
    style_table_header(ws_articles, 7)
    auto_fit_columns(ws_articles)

    ws_books = workbook.create_sheet('Libros')
    ws_books.append(['Título', 'Año', 'ISBN', 'Editorial', 'No. Investigadores'])
    for row in data['products']['books']:
        ws_books.append([
            row['title'],
            row.get('year'),
            row.get('isbn'),
            row.get('publisher'),
            row.get('researchersCount'),
        ])
    style_table_header(ws_books, 5)
    auto_fit_columns(ws_books)

    ws_chapters = workbook.create_sheet('Capitulos')
    ws_chapters.append(['Capítulo', 'Libro', 'Año', 'ISBN', 'Editorial', 'No. Investigadores'])
    for row in data['products']['bookChapters']:
        ws_chapters.append([
            row.get('chapterTitle'),
            row.get('bookTitle'),
            row.get('year'),
            row.get('isbn'),
            row.get('publisher'),
            row.get('researchersCount'),
        ])
    style_table_header(ws_chapters, 6)
    auto_fit_columns(ws_chapters)

    ws_theses = workbook.create_sheet('Tesis')
    ws_theses.append(['Rol', 'Título', 'Año', 'Tipo', 'Institución', 'Relacionados'])

    for row in data['theses']['asDirector']:
        ws_theses.append(['Director', row.get('title'), row.get('year'), row.get('thesisType'), row.get('institution'), ', '.join(row.get('students', []))])
    for row in data['theses']['asCotutor']:
        ws_theses.append(['Cotutor', row.get('title'), row.get('year'), row.get('thesisType'), row.get('institution'), ', '.join(row.get('students', []))])
    for row in data['theses']['asStudent']:
        ws_theses.append(['Estudiante', row.get('title'), row.get('year'), row.get('thesisType'), row.get('institution'), ', '.join(row.get('tutors', []))])
    style_table_header(ws_theses, 6)
    auto_fit_columns(ws_theses)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    safe_code = re.sub(r'[\\/:*?"<>|]+', '-', code_rh).strip() or 'SIN-CODIGO'
    safe_name = re.sub(r'[\\/:*?"<>|]+', ' ', researcher_name).strip() or 'Investigador'
    safe_name = re.sub(r'\s+', ' ', safe_name)
    filename = f"Ficha-CodigoRH-{safe_code}-{safe_name}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@api_view(['GET'])
def get_article_detail(request, article_id):
    article = Article.objects.select_related('group', 'journal', 'country', 'city').filter(id=article_id).first()
    if not article:
        return Response({'error': 'Article not found'}, status=404)

    authors = ArticleAuthor.objects.filter(article=article).select_related('researcher').order_by('order', 'author_name')
    related_articles = Article.objects.filter(group=article.group).exclude(id=article.id).distinct().order_by('-year', 'title')[:15]

    return Response(
        {
            'id': str(article.id),
            'title': article.title,
            'year': article.year,
            'doi': article.doi,
            'doiSuffix': article.doi_suffix,
            'issn': article.issn,
            'volume': article.volume,
            'issue': article.issue,
            'pages': article.pages,
            'productType': article.product_type,
            'group': {
                'id': str(article.group.id),
                'name': article.group.name,
                'code': article.group.code,
            } if article.group else None,
            'journal': {
                'id': str(article.journal.id),
                'name': article.journal.name,
                'issn': article.journal.issn,
            } if article.journal else None,
            'location': {
                'country': article.country.name if article.country else None,
                'city': article.city.name if article.city else None,
            },
            'authors': [
                {
                    'name': author.author_name,
                    'isGroupMember': author.is_group_member,
                    'researcherId': str(author.researcher.id) if author.researcher else None,
                    'researcherPath': f'/investigadores/{author.researcher.id}' if author.researcher else None,
                }
                for author in authors
            ],
            'related': {
                'articlesFromGroup': [
                    {
                        'id': str(item.id),
                        'title': item.title,
                        'year': item.year,
                        'detailPath': f'/articulos/{item.id}',
                    }
                    for item in related_articles
                ]
            }
        }
    )


@api_view(['GET'])
def get_book_detail(request, book_id):
    book = Book.objects.select_related('group', 'publisher', 'country', 'city').filter(id=book_id).first()
    if not book:
        return Response({'error': 'Book not found'}, status=404)

    authors = BookAuthor.objects.filter(book=book).select_related('researcher').order_by('order', 'author_name')
    related_chapters_qs = BookChapter.objects.filter(group=book.group).filter(
        Q(isbn=book.isbn) |
        Q(book_title__iexact=book.title)
    ).distinct().order_by('-year', 'chapter_title')

    related_books = Book.objects.filter(group=book.group).exclude(id=book.id).distinct().order_by('-year', 'title')[:15]

    return Response(
        {
            'id': str(book.id),
            'title': book.title,
            'year': book.year,
            'isbn': book.isbn,
            'productType': book.product_type,
            'group': {
                'id': str(book.group.id),
                'name': book.group.name,
                'code': book.group.code,
            } if book.group else None,
            'publisher': {
                'id': str(book.publisher.id),
                'name': book.publisher.name,
            } if book.publisher else None,
            'location': {
                'country': book.country.name if book.country else None,
                'city': book.city.name if book.city else None,
            },
            'authors': [
                {
                    'name': author.author_name,
                    'isGroupMember': author.is_group_member,
                    'researcherId': str(author.researcher.id) if author.researcher else None,
                    'researcherPath': f'/investigadores/{author.researcher.id}' if author.researcher else None,
                }
                for author in authors
            ],
            'related': {
                'chapters': [
                    {
                        'id': str(chapter.id),
                        'chapterTitle': chapter.chapter_title,
                        'bookTitle': chapter.book_title,
                        'year': chapter.year,
                        'detailPath': f'/capitulos/{chapter.id}',
                    }
                    for chapter in related_chapters_qs[:30]
                ],
                'booksFromGroup': [
                    {
                        'id': str(item.id),
                        'title': item.title,
                        'year': item.year,
                        'detailPath': f'/libros/{item.id}',
                    }
                    for item in related_books
                ],
            },
        }
    )


@api_view(['GET'])
def get_book_chapter_detail(request, chapter_id):
    chapter = BookChapter.objects.select_related('group', 'publisher').filter(id=chapter_id).first()
    if not chapter:
        return Response({'error': 'Book chapter not found'}, status=404)

    authors = ChapterAuthor.objects.filter(chapter=chapter).select_related('researcher').order_by('order', 'author_name')
    related_books_qs = Book.objects.filter(group=chapter.group).filter(
        Q(isbn=chapter.isbn) |
        Q(title__iexact=chapter.book_title)
    ).distinct().order_by('-year', 'title')

    related_chapters_qs = BookChapter.objects.filter(group=chapter.group).exclude(id=chapter.id).filter(
        Q(isbn=chapter.isbn) |
        Q(book_title__iexact=chapter.book_title)
    ).distinct().order_by('-year', 'chapter_title')

    return Response(
        {
            'id': str(chapter.id),
            'chapterTitle': chapter.chapter_title,
            'bookTitle': chapter.book_title,
            'year': chapter.year,
            'isbn': chapter.isbn,
            'productType': chapter.product_type,
            'group': {
                'id': str(chapter.group.id),
                'name': chapter.group.name,
                'code': chapter.group.code,
            } if chapter.group else None,
            'publisher': {
                'id': str(chapter.publisher.id),
                'name': chapter.publisher.name,
            } if chapter.publisher else None,
            'authors': [
                {
                    'name': author.author_name,
                    'isGroupMember': author.is_group_member,
                    'researcherId': str(author.researcher.id) if author.researcher else None,
                    'researcherPath': f'/investigadores/{author.researcher.id}' if author.researcher else None,
                }
                for author in authors
            ],
            'related': {
                'books': [
                    {
                        'id': str(book.id),
                        'title': book.title,
                        'year': book.year,
                        'detailPath': f'/libros/{book.id}',
                    }
                    for book in related_books_qs[:20]
                ],
                'chaptersFromSameBook': [
                    {
                        'id': str(item.id),
                        'chapterTitle': item.chapter_title,
                        'year': item.year,
                        'detailPath': f'/capitulos/{item.id}',
                    }
                    for item in related_chapters_qs[:20]
                ],
            },
        }
    )


def _build_group_detail_payload(group):
    # Researchers associated
    # group.members relates to GroupMember
    members = group.members.select_related('researcher').all()
    member_researcher_ids = {m.researcher_id for m in members if m.researcher_id}
    researchers_data = [
        {
            'id': str(m.researcher.id),
            'name': m.researcher.name,
            'codeRh': m.researcher.code_rh,
            'category': m.researcher.category,
            'membershipType': m.membership_type,
            'status': m.status,
            'detailPath': f'/investigadores/{m.researcher.id}',
        }
        for m in members
        if m.researcher
    ]

    # Articles
    articles = Article.objects.filter(group=group).order_by('-year', 'title')[:200]
    
    article_ids = [a.id for a in articles]
    
    article_categories_by_id = {}
    article_categories_qs = (
        ArticleCategory.objects.filter(article_id__in=article_ids, source__name__iexact='PUBLINDEX')
        .select_related('source')
        .values('article_id', 'year', 'category', 'source__name')
    )
    for row in article_categories_qs:
        article_categories_by_id.setdefault(row['article_id'], []).append(row)

    def resolve_article_category(article):
        rows = article_categories_by_id.get(article.id, [])
        if not rows:
            return {
                'category': 'N/R',
                'source': None,
            }

        target_year = article.year
        if target_year is not None:
            filtered = [r for r in rows if r.get('year') == target_year]
            if filtered:
                rows = filtered

        publindex_rows = [row for row in rows if row.get('category')]
        if publindex_rows:
            selected = max(publindex_rows, key=lambda row: row.get('year') or 0)
            return {
                'category': str(selected.get('category')),
                'source': 'Publindex',
            }

        return {
            'category': 'N/R',
            'source': None,
        }

        return {
            'category': 'N/R',
            'source': None,
        }
    
    # Calculate researchers explicitly grouped by product, getting names instead of just counts
    article_authors_qs = ArticleAuthor.objects.filter(article_id__in=article_ids).values('article_id', 'author_name', 'researcher_id').order_by('order')
    article_authors_by_id = {}
    for aa in article_authors_qs:
        article_authors_by_id.setdefault(aa['article_id'], []).append({
            'name': aa['author_name'],
            'isGroupMember': aa['researcher_id'] in member_researcher_ids if aa['researcher_id'] else False
        })

    articles_data = [
        {
            'id': str(a.id),
            'title': a.title,
            'year': a.year,
            'category': resolve_article_category(a)['category'],
            'categorySource': resolve_article_category(a)['source'],
            'doi': a.doi,
            'issn': a.issn,
            'authors': article_authors_by_id.get(a.id, []),
            'detailPath': f'/articulos/{a.id}'
        }
        for a in articles
    ]

    # Books
    books = Book.objects.filter(group=group).order_by('-year', 'title')[:200]
    book_ids = [b.id for b in books]
    book_authors_qs = BookAuthor.objects.filter(book_id__in=book_ids).values('book_id', 'author_name', 'researcher_id').order_by('order')
    book_authors_by_id = {}
    for ba in book_authors_qs:
        book_authors_by_id.setdefault(ba['book_id'], []).append({
            'name': ba['author_name'],
            'isGroupMember': ba['researcher_id'] in member_researcher_ids if ba['researcher_id'] else False
        })

    books_data = [
        {
            'id': str(b.id),
            'title': b.title,
            'year': b.year,
            'isbn': b.isbn,
            'authors': book_authors_by_id.get(b.id, []),
            'detailPath': f'/libros/{b.id}'
        }
        for b in books
    ]

    # Chapters
    chapters = BookChapter.objects.filter(group=group).order_by('-year', 'chapter_title')[:200]
    chapter_ids = [c.id for c in chapters]
    chapter_authors_qs = ChapterAuthor.objects.filter(chapter_id__in=chapter_ids).values('chapter_id', 'author_name', 'researcher_id').order_by('order')
    chapter_authors_by_id = {}
    for ca in chapter_authors_qs:
        chapter_authors_by_id.setdefault(ca['chapter_id'], []).append({
            'name': ca['author_name'],
            'isGroupMember': ca['researcher_id'] in member_researcher_ids if ca['researcher_id'] else False
        })

    chapters_data = [
        {
            'id': str(c.id),
            'chapterTitle': c.chapter_title,
            'bookTitle': c.book_title,
            'year': c.year,
            'isbn': c.isbn,
            'authors': chapter_authors_by_id.get(c.id, []),
            'detailPath': f'/capitulos/{c.id}'
        }
        for c in chapters
    ]

    # Theses
    theses = Thesis.objects.filter(group=group).order_by('-year', 'title')[:200]
    thesis_ids = [t.id for t in theses]
    thesis_authors_by_id = {}

    tutors_qs = ThesisTutor.objects.filter(thesis_id__in=thesis_ids).values('thesis_id', 'tutor_name', 'researcher_id').order_by('order')
    for tq in tutors_qs:
        thesis_authors_by_id.setdefault(tq['thesis_id'], []).append({
            'name': tq['tutor_name'],
            'isGroupMember': tq['researcher_id'] in member_researcher_ids if tq['researcher_id'] else False,
            'role': 'Tutor/Director'
        })

    students_qs = ThesisStudent.objects.filter(thesis_id__in=thesis_ids).values('thesis_id', 'student_name', 'researcher_id').order_by('order')
    for sq in students_qs:
        thesis_authors_by_id.setdefault(sq['thesis_id'], []).append({
            'name': sq['student_name'],
            'isGroupMember': sq['researcher_id'] in member_researcher_ids if sq['researcher_id'] else False,
            'role': 'Estudiante'
        })

    theses_data = [
        {
            'id': str(t.id),
            'title': t.title,
            'year': t.year,
            'thesisType': t.thesis_type,
            'institution': t.institution_obj.name if t.institution_obj else t.institution,
            'authors': thesis_authors_by_id.get(t.id, []),
            # we skip detailPath since we typically don't have thesis detail pages yet, but we could
        }
        for t in theses
    ]

    # Scientific events
    events = ScientificEvent.objects.filter(group=group).order_by('-start_date', 'title')[:200]
    event_ids = [event.id for event in events]
    institutions_by_event_id = {}
    institutions_qs = EventInstitution.objects.filter(event_id__in=event_ids).values('event_id', 'institution_name').order_by('id')
    for institution in institutions_qs:
        if not institution['institution_name']:
            continue
        institutions_by_event_id.setdefault(institution['event_id'], []).append(institution['institution_name'])

    events_data = [
        {
            'id': str(event.id),
            'title': event.title,
            'eventType': event.event_type,
            'scope': event.scope,
            'participationType': event.participation_type,
            'city': event.city_obj.name if event.city_obj else event.city,
            'country': event.country_obj.name if event.country_obj else None,
            'startDate': serialize_optional_datetime(event.start_date),
            'endDate': serialize_optional_datetime(event.end_date),
            'institutions': institutions_by_event_id.get(event.id, []),
        }
        for event in events
    ]
    
    # Generic Products
    generic_products = GenericProduct.objects.filter(group=group).order_by('-year', 'title')[:200]
    generic_ids = [gp.id for gp in generic_products]
    generic_authors_by_id = {}
    
    generic_authors_qs = GenericProductAuthor.objects.filter(product_id__in=generic_ids).values('product_id', 'name').order_by('id')
    for ga in generic_authors_qs:
        generic_authors_by_id.setdefault(ga['product_id'], []).append({
            'name': ga['name'],
            'isGroupMember': False  # GenericProductAuthor doesn't have researcher yet
        })
        
    generic_products_data = [
        {
            'id': str(gp.id),
            'title': gp.title,
            'year': gp.year,
            'category': gp.category,
            'tableName': gp.table_name,
            'authors': generic_authors_by_id.get(gp.id, []),
            'extraData': gp.extra_data,
        }
        for gp in generic_products
    ]

    return {
        'id': str(group.id),
        'code': group.code,
        'name': group.name,
        'leader': group.leader,
        'gruplacUrl': group.gruplac_url,
        'category': group.category,
        'formationDateInfo': group.formation_date_info,
        'department': group.department,
        'city': group.city,
        'certificationStatus': group.certification_status,
        'website': group.website,
        'email': group.email,
        'classificationValidity': group.classification_validity,
        'researchers': researchers_data,
        'products': {
            'articles': articles_data,
            'books': books_data,
            'bookChapters': chapters_data,
            'theses': theses_data,
            'events': events_data,
            'generic': generic_products_data,
        }
    }


def _build_typologies_payload(from_year=None, to_year=None):
    articles = list(Article.objects.select_related('group').order_by('-year', 'title'))
    article_ids = [article.id for article in articles]

    article_categories_by_id = {}
    article_categories_qs = (
        ArticleCategory.objects.filter(article_id__in=article_ids)
        .select_related('source')
        .values('article_id', 'year', 'category', 'source__name')
    )
    for row in article_categories_qs:
        article_categories_by_id.setdefault(row['article_id'], []).append(row)

    def resolve_article_category(article):
        rows = article_categories_by_id.get(article.id, [])
        if not rows:
            return {'category': 'N/R', 'source': None}

        target_year = article.year
        if target_year is not None:
            same_year_rows = [row for row in rows if row.get('year') == target_year]
            if same_year_rows:
                rows = same_year_rows

        publindex_rows = [row for row in rows if row.get('category')]
        if publindex_rows:
            selected = max(publindex_rows, key=lambda row: row.get('year') or 0)
            return {'category': str(selected['category']), 'source': 'Publindex'}
        return {'category': 'N/R', 'source': None}

    article_authors_by_id = {}
    article_authors_qs = ArticleAuthor.objects.filter(article_id__in=article_ids).values('article_id', 'author_name', 'researcher_id').order_by('order')
    for row in article_authors_qs:
        article_authors_by_id.setdefault(row['article_id'], []).append({
            'name': row['author_name'],
            'researcherPath': f"/investigadores/{row['researcher_id']}" if row.get('researcher_id') else None,
        })

    articles_data = []
    for article in articles:
        resolved_category = resolve_article_category(article)
        articles_data.append({
            'id': str(article.id),
            'title': article.title,
            'year': article.year,
            'category': resolved_category['category'],
            'categorySource': resolved_category['source'],
            'doi': article.doi,
            'issn': article.issn,
            'authors': article_authors_by_id.get(article.id, []),
            'detailPath': f'/articulos/{article.id}',
            'group': {
                'id': str(article.group.id),
                'name': article.group.name,
                'code': article.group.code,
                'category': article.group.category,
                'detailPath': f'/groups/{article.group.id}',
            } if article.group else None,
        })

    books = list(Book.objects.select_related('group').order_by('-year', 'title'))
    book_ids = [book.id for book in books]
    book_authors_by_id = {}
    book_authors_qs = BookAuthor.objects.filter(book_id__in=book_ids).values('book_id', 'author_name', 'researcher_id').order_by('order')
    for row in book_authors_qs:
        book_authors_by_id.setdefault(row['book_id'], []).append({
            'name': row['author_name'],
            'researcherPath': f"/investigadores/{row['researcher_id']}" if row.get('researcher_id') else None,
        })

    books_data = [
        {
            'id': str(book.id),
            'title': book.title,
            'year': book.year,
            'isbn': book.isbn,
            'authors': book_authors_by_id.get(book.id, []),
            'detailPath': f'/libros/{book.id}',
            'group': {
                'id': str(book.group.id),
                'name': book.group.name,
                'code': book.group.code,
                'category': book.group.category,
                'detailPath': f'/groups/{book.group.id}',
            } if book.group else None,
        }
        for book in books
    ]

    chapters = list(BookChapter.objects.select_related('group').order_by('-year', 'chapter_title'))
    chapter_ids = [chapter.id for chapter in chapters]
    chapter_authors_by_id = {}
    chapter_authors_qs = ChapterAuthor.objects.filter(chapter_id__in=chapter_ids).values('chapter_id', 'author_name', 'researcher_id').order_by('order')
    for row in chapter_authors_qs:
        chapter_authors_by_id.setdefault(row['chapter_id'], []).append({
            'name': row['author_name'],
            'researcherPath': f"/investigadores/{row['researcher_id']}" if row.get('researcher_id') else None,
        })

    chapters_data = [
        {
            'id': str(chapter.id),
            'chapterTitle': chapter.chapter_title,
            'bookTitle': chapter.book_title,
            'year': chapter.year,
            'isbn': chapter.isbn,
            'authors': chapter_authors_by_id.get(chapter.id, []),
            'detailPath': f'/capitulos/{chapter.id}',
            'group': {
                'id': str(chapter.group.id),
                'name': chapter.group.name,
                'code': chapter.group.code,
                'category': chapter.group.category,
                'detailPath': f'/groups/{chapter.group.id}',
            } if chapter.group else None,
        }
        for chapter in chapters
    ]

    theses = list(Thesis.objects.select_related('group', 'institution_obj').order_by('-year', 'title'))
    thesis_ids = [thesis.id for thesis in theses]
    thesis_authors_by_id = {}
    tutors_qs = ThesisTutor.objects.filter(thesis_id__in=thesis_ids).values('thesis_id', 'tutor_name', 'researcher_id').order_by('order')
    for row in tutors_qs:
        thesis_authors_by_id.setdefault(row['thesis_id'], []).append({
            'name': row['tutor_name'],
            'role': 'Tutor/Director',
            'researcherPath': f"/investigadores/{row['researcher_id']}" if row.get('researcher_id') else None,
        })
    students_qs = ThesisStudent.objects.filter(thesis_id__in=thesis_ids).values('thesis_id', 'student_name', 'researcher_id').order_by('order')
    for row in students_qs:
        thesis_authors_by_id.setdefault(row['thesis_id'], []).append({
            'name': row['student_name'],
            'role': 'Estudiante',
            'researcherPath': f"/investigadores/{row['researcher_id']}" if row.get('researcher_id') else None,
        })

    theses_data = [
        {
            'id': str(thesis.id),
            'title': thesis.title,
            'year': thesis.year,
            'thesisType': thesis.thesis_type,
            'institution': thesis.institution_obj.name if thesis.institution_obj else thesis.institution,
            'authors': thesis_authors_by_id.get(thesis.id, []),
            'group': {
                'id': str(thesis.group.id),
                'name': thesis.group.name,
                'code': thesis.group.code,
                'category': thesis.group.category,
                'detailPath': f'/groups/{thesis.group.id}',
            } if thesis.group else None,
        }
        for thesis in theses
    ]

    events = list(ScientificEvent.objects.select_related('group', 'city_obj', 'country_obj').order_by('-start_date', 'title'))
    event_ids = [event.id for event in events]
    institutions_by_event_id = {}
    institutions_qs = EventInstitution.objects.filter(event_id__in=event_ids).values('event_id', 'institution_name').order_by('id')
    for row in institutions_qs:
        if not row.get('institution_name'):
            continue
        institutions_by_event_id.setdefault(row['event_id'], []).append(row['institution_name'])

    events_data = [
        {
            'id': str(event.id),
            'title': event.title,
            'eventType': event.event_type,
            'scope': event.scope,
            'participationType': event.participation_type,
            'city': event.city_obj.name if event.city_obj else event.city,
            'country': event.country_obj.name if event.country_obj else None,
            'startDate': serialize_optional_datetime(event.start_date),
            'endDate': serialize_optional_datetime(event.end_date),
            'institutions': institutions_by_event_id.get(event.id, []),
            'group': {
                'id': str(event.group.id),
                'name': event.group.name,
                'code': event.group.code,
                'category': event.group.category,
                'detailPath': f'/groups/{event.group.id}',
            } if event.group else None,
        }
        for event in events
    ]

    generic_products = list(GenericProduct.objects.select_related('group').order_by('-year', 'title'))
    generic_ids = [product.id for product in generic_products]
    generic_authors_by_id = {}
    generic_authors_qs = GenericProductAuthor.objects.filter(product_id__in=generic_ids).values('product_id', 'name').order_by('id')
    for row in generic_authors_qs:
        generic_authors_by_id.setdefault(row['product_id'], []).append({'name': row['name']})

    generic_products_data = [
        {
            'id': str(product.id),
            'title': product.title,
            'year': product.year,
            'category': product.category,
            'tableName': product.table_name,
            'authors': generic_authors_by_id.get(product.id, []),
            'extraData': product.extra_data,
            'group': {
                'id': str(product.group.id),
                'name': product.group.name,
                'code': product.group.code,
                'category': product.group.category,
                'detailPath': f'/groups/{product.group.id}',
            } if product.group else None,
        }
        for product in generic_products
    ]

    products = {
        'articles': articles_data,
        'books': books_data,
        'bookChapters': chapters_data,
        'theses': theses_data,
        'events': events_data,
        'generic': generic_products_data,
    }

    return {
        'availableYears': get_available_years_from_products(products),
        'periodLabel': describe_export_year_filter(from_year, to_year),
        'products': filter_products_payload_by_year_range(products, from_year, to_year),
    }


@api_view(['GET'])
def get_typologies_overview(request):
    from_year = parse_optional_year_param(request.GET.get('fromYear'))
    to_year = parse_optional_year_param(request.GET.get('toYear'))
    if from_year and to_year and from_year > to_year:
        to_year = None

    return Response(_build_typologies_payload(from_year, to_year))


@api_view(['GET'])
def export_typologies_excel(request):
    from_year = parse_optional_year_param(request.GET.get('fromYear'))
    to_year = parse_optional_year_param(request.GET.get('toYear'))
    if from_year and to_year and from_year > to_year:
        to_year = None

    data = _build_typologies_payload(from_year, to_year)
    products = data['products']

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    header_fill = PatternFill(start_color='184336', end_color='184336', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    title_font = Font(size=18, bold=True, color='184336')
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB'),
    )

    def style_table_header(sheet, total_columns):
        for col_idx in range(1, total_columns + 1):
            cell = sheet.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = thin_border
        sheet.freeze_panes = 'A2'

    def style_table_body(sheet):
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
                cell.border = thin_border

    def auto_fit_columns(sheet, min_width=12, max_width=60):
        for col_cells in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col_cells[0].column)
            for cell in col_cells:
                max_len = max(max_len, len('' if cell.value is None else str(cell.value)))
            sheet.column_dimensions[col_letter].width = max(min(max_len + 2, max_width), min_width)

    def build_table_sheet(sheet_name, headers, rows, index=None):
        sheet = workbook.create_sheet(sheet_name, index) if index is not None else workbook.create_sheet(sheet_name)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        style_table_header(sheet, len(headers))
        style_table_body(sheet)
        auto_fit_columns(sheet)
        return sheet

    def normalize_export_value(value, fallback='N/A'):
        if value is None:
            return fallback
        if isinstance(value, str):
            cleaned = value.strip()
            return cleaned if cleaned else fallback
        return value

    def join_authors(authors, include_roles=False):
        entries = []
        for author in authors or []:
            name = (author.get('name') or '').strip()
            role = (author.get('role') or '').strip()
            if not name and not role:
                continue
            entries.append(f"{name} - {role}".strip(' -') if include_roles and role else (name or role))
        return '\n'.join(entries) if entries else 'N/A'

    def count_by_label(items, label_getter, empty_label='Sin clasificar'):
        counts = {}
        for item in items:
            label = normalize_export_value(label_getter(item), empty_label)
            counts[label] = counts.get(label, 0) + 1
        return sorted(counts.items(), key=lambda entry: (-entry[1], entry[0]))

    def add_summary_links(sheet):
        for row in sheet.iter_rows(min_row=2):
            target_sheet = row[3].value
            if not target_sheet or target_sheet not in workbook.sheetnames:
                continue
            row[3].hyperlink = Hyperlink(
                ref=row[3].coordinate,
                location=f"{quote_sheetname(target_sheet)}!A1",
                display=str(row[3].value),
            )
            row[3].font = Font(color='0563C1', underline='single', bold=(row[1].value == 'Total tipología'))

    cover_sheet = workbook.create_sheet('Portada', 0)
    cover_sheet.merge_cells('A1:E1')
    cover_sheet['A1'] = 'Panorama General de Tipologías'
    cover_sheet['A1'].font = title_font
    cover_sheet['A1'].alignment = Alignment(horizontal='center', vertical='center')

    cover_sheet['A3'] = 'Alcance'
    cover_sheet['B3'] = 'Todos los grupos con producción registrada'
    cover_sheet['A4'] = 'Periodo exportado'
    cover_sheet['B4'] = data['periodLabel']
    cover_sheet['A5'] = 'Fecha de generación'
    cover_sheet['B5'] = timezone.localtime().strftime('%Y-%m-%d %H:%M')

    overview_rows = [
        ('Nuevo conocimiento', len(products['articles']) + len(products['books']) + len(products['bookChapters'])),
        ('Formación RH', len(products['theses'])),
        ('Eventos científicos', len(products.get('events', []))),
        ('DTeI', sum(1 for item in products.get('generic', []) if item.get('category') == 'DTeI')),
        ('PASC', sum(1 for item in products.get('generic', []) if item.get('category') == 'PASC')),
        ('DP', sum(1 for item in products.get('generic', []) if item.get('category') == 'DP')),
    ]
    cover_sheet['D3'] = 'Totales visibles'
    cover_sheet['D3'].font = Font(bold=True, color='184336')
    for row_index, (label, value) in enumerate(overview_rows, start=4):
        cover_sheet[f'D{row_index}'] = label
        cover_sheet[f'E{row_index}'] = value
        for cell in (cover_sheet[f'D{row_index}'], cover_sheet[f'E{row_index}']):
            cell.border = thin_border
        cover_sheet[f'D{row_index}'].font = Font(bold=True, color='184336')
        cover_sheet[f'D{row_index}'].fill = PatternFill(start_color='E6F1EE', end_color='E6F1EE', fill_type='solid')
        cover_sheet[f'D{row_index}'].alignment = Alignment(horizontal='left', vertical='center')
        cover_sheet[f'E{row_index}'].alignment = Alignment(horizontal='center', vertical='center')

    for row in range(3, 6):
        cover_sheet[f'A{row}'].font = Font(bold=True, color='184336')
        cover_sheet[f'A{row}'].fill = PatternFill(start_color='E6F1EE', end_color='E6F1EE', fill_type='solid')
        cover_sheet[f'A{row}'].alignment = Alignment(horizontal='left', vertical='center')
        cover_sheet[f'A{row}'].border = thin_border
        cover_sheet[f'B{row}'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        cover_sheet[f'B{row}'].border = thin_border

    cover_sheet.column_dimensions['A'].width = 24
    cover_sheet.column_dimensions['B'].width = 50
    cover_sheet.column_dimensions['D'].width = 28
    cover_sheet.column_dimensions['E'].width = 16

    summary_sheet_rows = [
        ['Nuevo conocimiento', 'Total tipología', len(products['articles']) + len(products['books']) + len(products['bookChapters']), 'Nuevo conocimiento'],
        ['Nuevo conocimiento', 'Artículos', len(products['articles']), 'Nuevo conocimiento'],
        ['Nuevo conocimiento', 'Libros', len(products['books']), 'Nuevo conocimiento'],
        ['Nuevo conocimiento', 'Capítulos de libro', len(products['bookChapters']), 'Nuevo conocimiento'],
        ['Formación RH', 'Total tipología', len(products['theses']), 'Formación RH'],
        *[['Formación RH', subtype, total, 'Formación RH'] for subtype, total in count_by_label(products['theses'], lambda item: item.get('thesisType'))],
        ['Eventos científicos', 'Total tipología', len(products.get('events', [])), 'Eventos científicos'],
        *[['Eventos científicos', subtype, total, 'Eventos científicos'] for subtype, total in count_by_label(products.get('events', []), lambda item: item.get('eventType'))],
    ]
    for category in ('DTeI', 'PASC', 'DP'):
        category_rows = [item for item in products.get('generic', []) if item.get('category') == category]
        summary_sheet_rows.append([category, 'Total tipología', len(category_rows), category])
        summary_sheet_rows.extend([
            [category, subtype, total, category]
            for subtype, total in count_by_label(category_rows, lambda item: item.get('tableName'), 'Otros productos')
        ])

    ws_summary = build_table_sheet(
        'Resumen',
        ['Tipología', 'Subtipo o corte', 'Total registros', 'Ir a hoja'],
        summary_sheet_rows,
        index=1,
    )
    for row in ws_summary.iter_rows(min_row=2):
        if row[1].value == 'Total tipología':
            for cell in row:
                cell.font = Font(bold=True, color='184336')
                cell.fill = PatternFill(start_color='F1F7F5', end_color='F1F7F5', fill_type='solid')

    new_knowledge_rows = []
    for row in products['articles']:
        new_knowledge_rows.append([
            'Artículo',
            normalize_export_value(row.get('title')),
            row.get('year') or 'N/A',
            normalize_export_value(row.get('group', {}).get('code') if row.get('group') else None),
            normalize_export_value(row.get('group', {}).get('name') if row.get('group') else None),
            normalize_export_value(row.get('category'), 'N/R'),
            normalize_export_value(row.get('categorySource'), 'N/R'),
            normalize_export_value(row.get('issn')),
            normalize_export_value(row.get('doi')),
            join_authors(row.get('authors')),
        ])
    for row in products['books']:
        new_knowledge_rows.append([
            'Libro',
            normalize_export_value(row.get('title')),
            row.get('year') or 'N/A',
            normalize_export_value(row.get('group', {}).get('code') if row.get('group') else None),
            normalize_export_value(row.get('group', {}).get('name') if row.get('group') else None),
            'N/A',
            'N/A',
            normalize_export_value(row.get('isbn')),
            'N/A',
            join_authors(row.get('authors')),
        ])
    for row in products['bookChapters']:
        new_knowledge_rows.append([
            'Capítulo de libro',
            normalize_export_value(row.get('chapterTitle')),
            row.get('year') or 'N/A',
            normalize_export_value(row.get('group', {}).get('code') if row.get('group') else None),
            normalize_export_value(row.get('group', {}).get('name') if row.get('group') else None),
            normalize_export_value(row.get('bookTitle')),
            'N/A',
            normalize_export_value(row.get('isbn')),
            'N/A',
            join_authors(row.get('authors')),
        ])
    if new_knowledge_rows:
        build_table_sheet(
            'Nuevo conocimiento',
            ['Tipo de registro', 'Título del producto', 'Año', 'Código del grupo', 'Grupo', 'Categoría o contenedor', 'Fuente de categoría', 'ISSN / ISBN', 'DOI', 'Autores'],
            new_knowledge_rows,
        )

    thesis_rows = [
        [
            normalize_export_value(row.get('title')),
            row.get('year') or 'N/A',
            normalize_export_value(row.get('thesisType')),
            normalize_export_value(row.get('institution')),
            normalize_export_value(row.get('group', {}).get('code') if row.get('group') else None),
            normalize_export_value(row.get('group', {}).get('name') if row.get('group') else None),
            join_authors(row.get('authors'), include_roles=True),
        ]
        for row in products['theses']
    ]
    if thesis_rows:
        build_table_sheet(
            'Formación RH',
            ['Trabajo dirigido', 'Año', 'Nivel o tipo', 'Institución', 'Código del grupo', 'Grupo', 'Participantes'],
            thesis_rows,
        )

    event_rows = [
        [
            normalize_export_value(row.get('title')),
            normalize_export_value(row.get('startDate')),
            normalize_export_value(row.get('endDate')),
            normalize_export_value(row.get('eventType')),
            normalize_export_value(row.get('scope')),
            normalize_export_value(row.get('participationType')),
            normalize_export_value(row.get('group', {}).get('code') if row.get('group') else None),
            normalize_export_value(row.get('group', {}).get('name') if row.get('group') else None),
            '\n'.join(row.get('institutions') or []) or 'N/A',
        ]
        for row in products.get('events', [])
    ]
    if event_rows:
        build_table_sheet(
            'Eventos científicos',
            ['Evento', 'Fecha inicial', 'Fecha final', 'Clase de evento', 'Ámbito', 'Participación', 'Código del grupo', 'Grupo', 'Instituciones asociadas'],
            event_rows,
        )

    def build_generic_sheet(category, sheet_name):
        category_rows = [row for row in products.get('generic', []) if row.get('category') == category]
        if not category_rows:
            return
        extra_keys = ordered_generic_extra_data_keys(category_rows)
        has_raw_text = any(isinstance(row.get('extraData'), dict) and (row.get('extraData') or {}).get('raw_text') for row in category_rows)
        headers = [
            'Subtipo',
            'Producto o resultado',
            'Año',
            'Código del grupo',
            'Grupo',
            'Autores',
            *[humanize_extra_data_key(key) for key in extra_keys],
            *(['Soporte textual complementario'] if has_raw_text else []),
        ]
        rows = []
        for row in category_rows:
            extra_data = row.get('extraData') or {}
            rows.append([
                normalize_export_value(row.get('tableName')),
                normalize_export_value(row.get('title'), 'S/T'),
                row.get('year') or 'N/A',
                normalize_export_value(row.get('group', {}).get('code') if row.get('group') else None),
                normalize_export_value(row.get('group', {}).get('name') if row.get('group') else None),
                join_authors(row.get('authors')),
                *[stringify_extra_data_value(extra_data.get(key)) for key in extra_keys],
                *([stringify_extra_data_value(extra_data.get('raw_text'))] if has_raw_text else []),
            ])
        build_table_sheet(sheet_name, headers, rows)

    build_generic_sheet('DTeI', 'DTeI')
    build_generic_sheet('PASC', 'PASC')
    build_generic_sheet('DP', 'DP')
    add_summary_links(ws_summary)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"Tipologias-UCEVA-{timezone.localtime().strftime('%Y%m%d-%H%M')}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response

@api_view(['GET'])
def get_groups_list(request):
    """
    Returns all research groups with summary data suitable for listing cards.
    Uses correlated subqueries instead of a multi-table JOIN to avoid
    Cartesian-product blowup that caused BufFileRead disk spills.
    """
    from .models import Article, Book, BookChapter, Thesis, GroupMember

    def _count_sq(model, fk='group_id'):
        return Coalesce(
            Subquery(
                model.objects
                    .filter(**{fk: OuterRef('pk')})
                    .values(fk)
                    .annotate(c=Count('id'))
                    .values('c')[:1],
                output_field=IntegerField(),
            ),
            0,
        )

    groups = (
        ResearchGroup.objects
        .only('id', 'name', 'code', 'category', 'leader')
        .annotate(
            articles_count=_count_sq(Article),
            books_count=_count_sq(Book),
            chapters_count=_count_sq(BookChapter),
            theses_count=_count_sq(Thesis),
            researchers_count=_count_sq(GroupMember),
        )
        .order_by('name')
    )
    data = [
        {
            'id': str(g.id),
            'name': g.name,
            'code': g.code,
            'category': g.category,
            'leader': g.leader,
            'articles_count': g.articles_count,
            'books_count': g.books_count,
            'chapters_count': g.chapters_count,
            'theses_count': g.theses_count,
            'researchers_count': g.researchers_count,
        }
        for g in groups
    ]
    return Response(data)


@api_view(['GET'])
def get_researchers_list(request):
    """
    Returns all researchers with production counters and derived active status
    for the global researchers listing page.
    """

    def _count_sq(model, fk='researcher_id', counted_field='id', active_only=False):
        queryset = model.objects.filter(**{fk: OuterRef('pk')})
        if active_only:
            queryset = queryset.exclude(status__iexact='INACTIVE')

        return Coalesce(
            Subquery(
                queryset
                    .values(fk)
                    .annotate(c=Count(counted_field, distinct=True))
                    .values('c')[:1],
                output_field=IntegerField(),
            ),
            0,
        )

    researchers = (
        Researcher.objects
        .only('id', 'name', 'code_rh', 'category', 'education_level', 'city', 'department', 'university', 'cvlac_url')
        .annotate(
            articles_count=_count_sq(ArticleAuthor, counted_field='article_id'),
            books_count=_count_sq(BookAuthor, counted_field='book_id'),
            chapters_count=_count_sq(ChapterAuthor, counted_field='chapter_id'),
            thesis_tutor_count=_count_sq(ThesisTutor, counted_field='thesis_id'),
            thesis_student_count=_count_sq(ThesisStudent, counted_field='thesis_id'),
            groups_count=_count_sq(GroupMember, counted_field='group_id'),
            active_groups_count=_count_sq(GroupMember, counted_field='group_id', active_only=True),
        )
        .order_by('name')
    )

    data = []
    for researcher in researchers:
        theses_count = (researcher.thesis_tutor_count or 0) + (researcher.thesis_student_count or 0)
        total_count = (
            (researcher.articles_count or 0)
            + (researcher.books_count or 0)
            + (researcher.chapters_count or 0)
            + theses_count
        )
        status = 'ACTIVE' if (researcher.active_groups_count or 0) > 0 else 'INACTIVE'

        data.append({
            'id': str(researcher.id),
            'name': researcher.name,
            'codeRh': researcher.code_rh,
            'category': researcher.category,
            'educationLevel': researcher.education_level,
            'city': researcher.city,
            'department': researcher.department,
            'university': researcher.university,
            'cvlacUrl': researcher.cvlac_url,
            'status': status,
            'articlesCount': researcher.articles_count or 0,
            'booksCount': researcher.books_count or 0,
            'chaptersCount': researcher.chapters_count or 0,
            'thesesCount': theses_count,
            'groupsCount': researcher.groups_count or 0,
            'totalCount': total_count,
            'detailPath': f'/researchers/{researcher.id}',
        })

    return Response(data)

@api_view(['GET'])
def get_group_detail(request, group_id):
    group = ResearchGroup.objects.filter(id=group_id).first()
    if not group:
        return Response({'error': 'Group not found'}, status=404)

    return Response(_build_group_detail_payload(group))

@api_view(['GET'])
def export_group_detail_excel(request, group_id):
    group = ResearchGroup.objects.filter(id=group_id).first()
    if not group:
        return Response({'error': 'Group not found'}, status=404)

    from_year = parse_optional_year_param(request.GET.get('fromYear'))
    to_year = parse_optional_year_param(request.GET.get('toYear'))
    if from_year and to_year and from_year > to_year:
        to_year = None

    data = filter_group_detail_payload_by_year_range(
        _build_group_detail_payload(group),
        from_year,
        to_year,
    )

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    # Frontend primary tone: hsl(162 47% 18%) ~= #184336
    header_fill = PatternFill(start_color='184336', end_color='184336', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    title_font = Font(size=18, bold=True, color='184336')
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB'),
    )

    def style_table_header(sheet, total_columns):
        for col_idx in range(1, total_columns + 1):
            cell = sheet.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = thin_border
        sheet.freeze_panes = 'A2'

    def auto_fit_columns(sheet, min_width=12, max_width=60):
        for col_cells in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col_cells[0].column)
            for cell in col_cells:
                cell_value = '' if cell.value is None else str(cell.value)
                max_len = max(max_len, len(cell_value))
            sheet.column_dimensions[col_letter].width = max(min(max_len + 2, max_width), min_width)

    def style_table_body(sheet):
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
                cell.border = thin_border

    def build_table_sheet(sheet_name, headers, rows, index=None):
        sheet = workbook.create_sheet(sheet_name, index) if index is not None else workbook.create_sheet(sheet_name)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        style_table_header(sheet, len(headers))
        style_table_body(sheet)
        auto_fit_columns(sheet)
        return sheet

    def join_authors(authors, only_group=None, include_roles=False):
        entries = []
        for author in authors or []:
            is_group_member = author.get('isGroupMember')
            if only_group is True and not is_group_member:
                continue
            if only_group is False and is_group_member:
                continue

            name = (author.get('name') or '').strip()
            role = (author.get('role') or '').strip()
            if not name and not role:
                continue

            if include_roles and role:
                entries.append(f"{name} - {role}".strip(' -'))
            else:
                entries.append(name or role)

        return '\n'.join(entries) if entries else 'N/A'

    def normalize_export_value(value, fallback='N/A'):
        if value is None:
            return fallback
        if isinstance(value, str):
            cleaned = value.strip()
            return cleaned if cleaned else fallback
        return value

    def count_by_label(items, label_getter, empty_label='Sin clasificar'):
        counts = {}
        for item in items:
            raw_label = label_getter(item)
            label = normalize_export_value(raw_label, empty_label)
            counts[label] = counts.get(label, 0) + 1
        return sorted(counts.items(), key=lambda entry: (-entry[1], entry[0]))

    def add_summary_links(sheet):
        for row in sheet.iter_rows(min_row=2):
            target_sheet = row[3].value
            if not target_sheet or target_sheet not in workbook.sheetnames:
                continue
            row[3].hyperlink = Hyperlink(
                ref=row[3].coordinate,
                location=f"{quote_sheetname(target_sheet)}!A1",
                display=str(row[3].value),
            )
            row[3].font = Font(
                color='0563C1',
                underline='single',
                bold=(row[1].value == 'Total tipología' or row[0].value == 'Investigadores'),
            )

    cover_sheet = workbook.create_sheet('Portada', 0)
    group_code = (group.code or 'SIN-CODIGO').strip()
    group_name = (group.name or 'Grupo').strip()

    cover_sheet.merge_cells('A1:E1')
    cover_sheet['A1'] = 'Ficha de Grupo de Investigación'
    cover_sheet['A1'].font = title_font
    cover_sheet['A1'].alignment = Alignment(horizontal='center', vertical='center')

    cover_sheet['A3'] = 'Código MinCiencias'
    cover_sheet['B3'] = group_code
    cover_sheet['A4'] = 'Nombre del Grupo'
    cover_sheet['B4'] = group_name
    cover_sheet['A5'] = 'Categoría'
    cover_sheet['B5'] = group.category or 'N/A'
    cover_sheet['A6'] = 'Líder'
    cover_sheet['B6'] = group.leader or 'N/A'
    cover_sheet['A7'] = 'Sitio web'
    cover_sheet['B7'] = normalize_export_value(group.website)
    cover_sheet['A8'] = 'Gruplac / MinCiencias'
    cover_sheet['B8'] = normalize_export_value(data.get('gruplacUrl'))
    cover_sheet['A9'] = 'Periodo exportado'
    cover_sheet['B9'] = describe_export_year_filter(from_year, to_year)
    cover_sheet['A10'] = 'Fecha de generación'
    cover_sheet['B10'] = timezone.localtime().strftime('%Y-%m-%d %H:%M')

    summary_rows = [
        ('Investigadores activos', sum(1 for researcher in data['researchers'] if researcher.get('status') != 'INACTIVE')),
        ('Investigadores inactivos', sum(1 for researcher in data['researchers'] if researcher.get('status') == 'INACTIVE')),
        ('Nuevo conocimiento', len(data['products']['articles']) + len(data['products']['books']) + len(data['products']['bookChapters'])),
        ('Formación de recurso humano', len(data['products']['theses'])),
        ('Eventos científicos', len(data['products'].get('events', []))),
        ('DTeI', sum(1 for product in data['products'].get('generic', []) if product.get('category') == 'DTeI')),
        ('PASC', sum(1 for product in data['products'].get('generic', []) if product.get('category') == 'PASC')),
        ('DP', sum(1 for product in data['products'].get('generic', []) if product.get('category') == 'DP')),
    ]

    cover_sheet['D3'] = 'Resumen exportado'
    cover_sheet['D3'].font = Font(bold=True, color='184336')
    for row_index, (label, value) in enumerate(summary_rows, start=4):
        cover_sheet[f'D{row_index}'] = label
        cover_sheet[f'E{row_index}'] = value
        label_cell = cover_sheet[f'D{row_index}']
        value_cell = cover_sheet[f'E{row_index}']
        label_cell.font = Font(bold=True, color='184336')
        label_cell.fill = PatternFill(start_color='E6F1EE', end_color='E6F1EE', fill_type='solid')
        label_cell.alignment = Alignment(horizontal='left', vertical='center')
        label_cell.border = thin_border
        value_cell.alignment = Alignment(horizontal='center', vertical='center')
        value_cell.border = thin_border

    for row in range(3, 11):
        label_cell = cover_sheet[f'A{row}']
        value_cell = cover_sheet[f'B{row}']
        label_cell.font = Font(bold=True, color='184336')
        label_cell.fill = PatternFill(start_color='E6F1EE', end_color='E6F1EE', fill_type='solid')
        label_cell.alignment = Alignment(horizontal='left', vertical='center')
        label_cell.border = thin_border
        value_cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        value_cell.border = thin_border

    cover_sheet.column_dimensions['A'].width = 28
    cover_sheet.column_dimensions['B'].width = 60
    cover_sheet.column_dimensions['D'].width = 30
    cover_sheet.column_dimensions['E'].width = 16
    cover_sheet.row_dimensions[1].height = 30

    summary_sheet_rows = [
        ['Investigadores', 'Total registrados', len(data['researchers']), 'Investigadores'],
        ['Investigadores', 'Activos', sum(1 for researcher in data['researchers'] if researcher.get('status') != 'INACTIVE'), 'Investigadores'],
        ['Investigadores', 'Inactivos', sum(1 for researcher in data['researchers'] if researcher.get('status') == 'INACTIVE'), 'Investigadores'],
        ['Nuevo conocimiento', 'Total tipología', len(data['products']['articles']) + len(data['products']['books']) + len(data['products']['bookChapters']), 'Nuevo conocimiento'],
        ['Nuevo conocimiento', 'Artículos', len(data['products']['articles']), 'Nuevo conocimiento'],
        ['Nuevo conocimiento', 'Libros', len(data['products']['books']), 'Nuevo conocimiento'],
        ['Nuevo conocimiento', 'Capítulos de libro', len(data['products']['bookChapters']), 'Nuevo conocimiento'],
        ['Formación RH', 'Total tipología', len(data['products']['theses']), 'Formación RH'],
    ]

    summary_sheet_rows.extend([
        ['Formación RH', subtype, total, 'Formación RH']
        for subtype, total in count_by_label(data['products']['theses'], lambda item: item.get('thesisType'))
    ])
    summary_sheet_rows.append(['Eventos científicos', 'Total tipología', len(data['products'].get('events', [])), 'Eventos científicos'])
    summary_sheet_rows.extend([
        ['Eventos científicos', subtype, total, 'Eventos científicos']
        for subtype, total in count_by_label(data['products'].get('events', []), lambda item: item.get('eventType'))
    ])

    for category in ('DTeI', 'PASC', 'DP'):
        category_rows = [row for row in data['products'].get('generic', []) if row.get('category') == category]
        summary_sheet_rows.append([category, 'Total tipología', len(category_rows), category])
        summary_sheet_rows.extend([
            [category, subtype, total, category]
            for subtype, total in count_by_label(category_rows, lambda item: item.get('tableName'), 'Otros productos')
        ])

    ws_summary = build_table_sheet(
        'Resumen',
        ['Tipología', 'Subtipo o corte', 'Total registros', 'Ir a hoja'],
        summary_sheet_rows,
        index=1,
    )
    for row in ws_summary.iter_rows(min_row=2):
        if row[1].value == 'Total tipología' or row[0].value == 'Investigadores':
            for cell in row:
                cell.font = Font(bold=True, color='184336')
                cell.fill = PatternFill(start_color='F1F7F5', end_color='F1F7F5', fill_type='solid')

    build_table_sheet(
        'Investigadores',
        ['Código RH', 'Investigador', 'Categoría', 'Vinculación', 'Estado'],
        [
            [
                row.get('codeRh') or 'N/A',
                row.get('name'),
                row.get('category') or 'N/A',
                row.get('membershipType') or 'N/A',
                'Activo' if row.get('status') == 'ACTIVE' else ('Inactivo' if row.get('status') == 'INACTIVE' else (row.get('status') or 'N/A')),
            ]
            for row in data['researchers']
        ]
    )

    new_knowledge_rows = []
    for row in data['products']['articles']:
        category_value = normalize_export_value(row.get('category'), 'N/R')
        source_value = normalize_export_value(row.get('categorySource'), 'N/R') if category_value != 'N/R' else 'N/R'
        new_knowledge_rows.append([
            'Artículo',
            normalize_export_value(row.get('title')),
            row.get('year') or 'N/A',
            'N/A',
            category_value,
            source_value,
            normalize_export_value(row.get('issn')),
            normalize_export_value(row.get('doi')),
            join_authors(row.get('authors'), only_group=True),
            join_authors(row.get('authors'), only_group=False),
        ])

    for row in data['products']['books']:
        new_knowledge_rows.append([
            'Libro',
            normalize_export_value(row.get('title')),
            row.get('year') or 'N/A',
            'N/A',
            'N/A',
            'N/A',
            normalize_export_value(row.get('isbn')),
            'N/A',
            join_authors(row.get('authors'), only_group=True),
            join_authors(row.get('authors'), only_group=False),
        ])

    for row in data['products']['bookChapters']:
        new_knowledge_rows.append([
            'Capítulo de libro',
            normalize_export_value(row.get('chapterTitle')),
            row.get('year') or 'N/A',
            normalize_export_value(row.get('bookTitle')),
            'N/A',
            'N/A',
            normalize_export_value(row.get('isbn')),
            'N/A',
            join_authors(row.get('authors'), only_group=True),
            join_authors(row.get('authors'), only_group=False),
        ])

    if new_knowledge_rows:
        build_table_sheet(
            'Nuevo conocimiento',
            ['Tipo de registro', 'Título del producto', 'Año', 'Libro o contenedor', 'Categoría', 'Fuente de categoría', 'ISSN / ISBN', 'DOI', 'Autores del grupo', 'Autores externos'],
            new_knowledge_rows,
        )

    theses_rows = [
        [
            normalize_export_value(row.get('title')),
            row.get('year') or 'N/A',
            normalize_export_value(row.get('thesisType')),
            normalize_export_value(row.get('institution')),
            join_authors(row.get('authors'), only_group=True, include_roles=True),
            join_authors(row.get('authors'), only_group=False, include_roles=True),
        ]
        for row in data['products']['theses']
    ]
    if theses_rows:
        build_table_sheet(
            'Formación RH',
            ['Trabajo dirigido', 'Año', 'Nivel o tipo', 'Institución', 'Participación del grupo', 'Participación externa'],
            theses_rows,
        )

    event_rows = [
        [
            normalize_export_value(row.get('title')),
            normalize_export_value(row.get('startDate')),
            normalize_export_value(row.get('endDate')),
            normalize_export_value(row.get('eventType')),
            normalize_export_value(row.get('scope')),
            normalize_export_value(row.get('participationType')),
            normalize_export_value(row.get('city')),
            normalize_export_value(row.get('country')),
            '\n'.join(row.get('institutions') or []) or 'N/A',
        ]
        for row in data['products'].get('events', [])
    ]
    if event_rows:
        build_table_sheet(
            'Eventos científicos',
            ['Evento', 'Fecha inicial', 'Fecha final', 'Clase de evento', 'Ámbito', 'Tipo de participación', 'Ciudad', 'País', 'Instituciones asociadas'],
            event_rows,
        )

    def build_generic_category_sheet(category, sheet_name):
        category_rows = [row for row in data['products'].get('generic', []) if row.get('category') == category]
        if not category_rows:
            return

        extra_keys = ordered_generic_extra_data_keys(category_rows)
        has_raw_text = any(
            isinstance(row.get('extraData'), dict) and (row.get('extraData') or {}).get('raw_text')
            for row in category_rows
        )
        headers = [
            'Subtipo',
            'Producto o resultado',
            'Año',
            'Autores del grupo',
            'Autores externos',
            *[humanize_extra_data_key(key) for key in extra_keys],
            *(['Soporte textual complementario'] if has_raw_text else []),
        ]
        rows = []
        for row in category_rows:
            extra_data = row.get('extraData') or {}
            rows.append([
                normalize_export_value(row.get('tableName')),
                normalize_export_value(row.get('title'), 'S/T'),
                row.get('year') or 'N/A',
                join_authors(row.get('authors'), only_group=True),
                join_authors(row.get('authors'), only_group=False),
                *[stringify_extra_data_value(extra_data.get(key)) for key in extra_keys],
                *([stringify_extra_data_value(extra_data.get('raw_text'))] if has_raw_text else []),
            ])
        build_table_sheet(sheet_name, headers, rows)

    build_generic_category_sheet('DTeI', 'DTeI')
    build_generic_category_sheet('PASC', 'PASC')
    build_generic_category_sheet('DP', 'DP')

    add_summary_links(ws_summary)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    safe_code = re.sub(r'[\\/:*?"<>|]+', '-', group_code).strip() or 'SIN-CODIGO'
    safe_name = re.sub(r'[\\/:*?"<>|]+', ' ', group_name).strip() or 'Grupo'
    safe_name = re.sub(r'\s+', ' ', safe_name)
    filename = f"Ficha-Grupo-{safe_code}-{safe_name}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response

@api_view(['GET'])
def get_stats(request):
    researchers_count = Researcher.objects.count()
    
    # Count unique productions based on Title + Year to avoid duplicates from multiple groups
    articles_count = Article.objects.values('title', 'year').distinct().count()
    theses_count = Thesis.objects.values('title', 'year').distinct().count()
    
    books_count = Book.objects.values('title', 'year').distinct().count()
    chapters_count = BookChapter.objects.values('chapter_title', 'book_title', 'year').distinct().count()
    books_chapters_count = books_count + chapters_count
    
    generic_count = GenericProduct.objects.values('title', 'year', 'table_name').distinct().count()

    events_count = ScientificEvent.objects.annotate(
        event_year=ExtractYear('start_date')
    ).values('title', 'event_year').distinct().count()
    
    # Thesis breakdown by type (using distinct title to avoid duplicates)
    thesis_qs = Thesis.objects.values('thesis_type').annotate(count=Count('title', distinct=True)).order_by('-count')
    thesis_breakdown = [
        {'label': t['thesis_type'] or 'Sin Clasificar', 'value': t['count']} 
        for t in thesis_qs
    ]

    # Researcher breakdown by category
    researcher_qs = Researcher.objects.values('category').annotate(count=Count('id')).order_by('-count')
    researcher_breakdown = [
        {'label': r['category'] or 'Sin Categoría', 'value': r['count']}
        for r in researcher_qs
    ]

    # Events breakdown (3 sections) - using distinct title for consistency
    events_by_type = ScientificEvent.objects.values('event_type').annotate(count=Count('title', distinct=True)).order_by('-count')[:5]
    events_by_scope = ScientificEvent.objects.values('scope').annotate(count=Count('title', distinct=True)).order_by('-count')
    
    # Process participation type manually to handle comma-separated values
    participation_counts = {}
    
    # Get distinct event participations to avoid double counting shared events
    # We select title/year/participation_type distinct combinations
    unique_participations = ScientificEvent.objects.annotate(
        event_year=ExtractYear('start_date')
    ).exclude(
        participation_type__isnull=True
    ).exclude(
        participation_type=''
    ).values('title', 'event_year', 'participation_type').distinct()
    
    for row in unique_participations:
        p_string = row['participation_type']
        if not p_string: continue
        # Split by comma
        parts = p_string.split(',')
        for part in parts:
            clean_part = part.strip()
            if not clean_part: continue
            # Normalize casing to Title Case to group "Ponente" and "ponente"
            clean_part = clean_part.title()
            participation_counts[clean_part] = participation_counts.get(clean_part, 0) + 1

    # Convert dict to sorted list
    participation_list = [
        {'participation_type': k, 'count': v} 
        for k, v in participation_counts.items()
    ]
    participation_list.sort(key=lambda x: x['count'], reverse=True)
    
    # Limit to top 10 for display
    participation_list = participation_list[:10]

    events_sections = [
        {
            "title": "Por Tipo",
            "items": [{'label': x['event_type'] or 'N/A', 'value': x['count']} for x in events_by_type]
        },
        {
            "title": "Por Ámbito",
            "items": [{'label': (x['scope'] if x['scope'] and x['scope'].lower() != 'null' else 'Otro'), 'value': x['count']} for x in events_by_scope]
        },
        {
            "title": "Por Participación",
            "items": [{'label': x['participation_type'] or 'N/A', 'value': x['count']} for x in participation_list]
        }
    ]

    data = [
        { 
            'label': 'Investigadores', 
            'value': str(researchers_count), 
            'breakdown': researcher_breakdown
        },
        { 'label': 'Artículos', 'value': str(articles_count) },
        {
            'label': 'Tesis de Grado',
            'value': str(theses_count),
            'breakdown': thesis_breakdown
        },
        { 'label': 'Libros & Capítulos', 'value': str(books_chapters_count) },
        { 'label': 'Otros Productos (DTeI, DP)', 'value': str(generic_count) },
        {
            'label': 'Eventos Científicos',
            'value': str(events_count),
            'sections': events_sections
        }
    ]
    
    return Response(data)

def _count_related_country_items(model, foreign_key_name):
    return Coalesce(
        Subquery(
            model.objects
                .filter(**{foreign_key_name: OuterRef('pk')})
                .values(foreign_key_name)
                .annotate(c=Count('id', distinct=True))
                .values('c')[:1],
            output_field=IntegerField(),
        ),
        0,
    )


def _get_country_network_records():
    countries_qs = Country.objects.annotate(
        articles_count=_count_related_country_items(Article, 'country_id'),
        books_count=_count_related_country_items(Book, 'country_id'),
        events_count=_count_related_country_items(ScientificEvent, 'country_obj_id'),
    ).order_by('name')

    grouped_records = {}
    for country in countries_qs:
        total_count = (country.articles_count or 0) + (country.books_count or 0) + (country.events_count or 0)
        if total_count <= 0:
            continue

        canonical_key = canonicalize_country_name(country.name)
        display_name = COUNTRY_CANONICAL_DISPLAY_NAMES.get(canonical_key, country.name.title())
        record = grouped_records.setdefault(canonical_key, {
            'id': f'country:{canonical_key}',
            'name': display_name,
            'name_clean': remove_accents(display_name).title(),
            'name_normalized': canonical_key,
            'articles_count': 0,
            'books_count': 0,
            'events_count': 0,
            'total_count': 0,
        })

        record['articles_count'] += country.articles_count or 0
        record['books_count'] += country.books_count or 0
        record['events_count'] += country.events_count or 0
        record['total_count'] += total_count

    return list(grouped_records.values())


def _build_country_network_summary(record, is_colombia=False):
    if is_colombia:
        return 'Sede principal de la producción nacional.'

    summary_parts = []
    if record['articles_count']:
        summary_parts.append(f"{record['articles_count']} artículos")
    if record['books_count']:
        summary_parts.append(f"{record['books_count']} libros")
    if record['events_count']:
        summary_parts.append(f"{record['events_count']} eventos")

    return ' / '.join(summary_parts) if summary_parts else 'Cooperación activa en investigación.'


@api_view(['GET'])
def get_cooperation_map(request):
    """
    Returns list of countries involved in georeferenced production, with coordinates.
    """
    hubs = []

    for record in _get_country_network_records():
        coords = COUNTRY_COORDINATES.get(record['name_normalized'])
        if not coords:
            continue

        is_colombia = record['name_normalized'] == 'colombia'

        hub = {
            'lat': coords['lat'],
            'lng': coords['lng'],
            'id': record['id'],
            'name': record['name'],
            'name_clean': record['name_clean'],
            'value': record['total_count'],
            'info': 'Nodo central UCEVA' if is_colombia else f"{record['total_count']} registros asociados",
            'summary': _build_country_network_summary(record, is_colombia=is_colombia),
            'link': f"/report/{record['name']}"
        }
        if is_colombia:
            hubs.insert(0, hub)
        else:
            hubs.append(hub)

    if not any(hub['name'].lower() == 'colombia' for hub in hubs):
        hubs.insert(0, {
            'lat': 4.5709,
            'lng': -74.2973,
            'id': 'colombia-main',
            'name': 'Colombia',
            'name_clean': 'Colombia',
            'value': 0,
            'info': 'Nodo central UCEVA',
            'summary': 'Sede principal de la producción nacional.',
            'link': '/report/colombia'
        })

    return Response(hubs)


@api_view(['GET'])
def get_map_metrics(request):
    records = _get_country_network_records()
    national_record = next((record for record in records if record['name_normalized'] == 'colombia'), None)
    national_total = national_record['total_count'] if national_record else 0
    international_total = sum(record['total_count'] for record in records if record['name_normalized'] != 'colombia')
    countries_in_network = len(records)
    featured_country = max(
        (record for record in records if record['name_normalized'] != 'colombia'),
        key=lambda record: (record['total_count'], record['name']),
        default=None,
    )

    return Response({
        'nationalProduction': national_total,
        'internationalProduction': international_total,
        'countriesInNetwork': countries_in_network,
        'featuredCountry': {
            'name': featured_country['name'],
            'summary': _build_country_network_summary(featured_country),
            'totalRecords': featured_country['total_count'],
            'link': f"/report/{featured_country['name']}",
        } if featured_country else None,
    })


@api_view(['GET'])
def get_country_reports_list(request):
    """
    Returns countries with associated production counts for the reports index page.
    """
    reports = []

    for record in sorted(
        _get_country_network_records(),
        key=lambda item: (-item['total_count'], item['name'])
    ):
        is_colombia = record['name_normalized'] == 'colombia'
        reports.append({
            'id': record['id'],
            'name': record['name'],
            'scope': 'Nacional' if is_colombia else 'Internacional',
            'summary': _build_country_network_summary(record, is_colombia=is_colombia),
            'articlesCount': record['articles_count'],
            'booksCount': record['books_count'],
            'eventsCount': record['events_count'],
            'totalCount': record['total_count'],
            'detailPath': f"/report/{record['name']}",
        })

    return Response(reports)

from scienti.models import Country, Researcher, Article, Thesis, Book, BookChapter, ResearchGroup

@api_view(['GET'])
def get_country_details(request, country_id):
    import urllib.parse
    
    # Ensure proper decoding of URL parameter
    country_id = urllib.parse.unquote(country_id)
    matched_countries, display_name, canonical_key = resolve_country_group(country_id)

    if not matched_countries:
        return Response({"error": f"Country '{country_id}' not found"}, status=404)

    country_ids = [country.id for country in matched_countries]
    countries_qs = Country.objects.filter(id__in=country_ids)

    # Calculate stats
    articles_qs = Article.objects.filter(country__in=countries_qs)
    books_qs = Book.objects.filter(country__in=countries_qs)
    
    # Try to find chapters via Book title match (best effort)
    # Get titles of books published in this country
    book_titles_in_country = books_qs.values_list('title', flat=True)
    chapters_qs = BookChapter.objects.filter(book_title__in=book_titles_in_country)
    
    articles_count = articles_qs.count()
    books_count = books_qs.count()
    chapters_count = chapters_qs.count()
    
    # Thesis doesn't have country
    thesis_count = 0 

    recent_articles = articles_qs.order_by('-year', '-title')[:3].values(
        'title', 'year', 'journal__name', 'doi'
    )
    
    recent_books = books_qs.order_by('-year', '-title')[:3].values(
        'title', 'year', 'isbn'
    )

    researcher_ids = Article.objects.filter(country__in=countries_qs).values_list('authors__researcher', flat=True)
    researcher_ids = [rid for rid in researcher_ids if rid]
    
    book_researcher_ids = Book.objects.filter(country__in=countries_qs).values_list('authors__researcher', flat=True)
    
    all_researcher_ids = list(researcher_ids) + list(book_researcher_ids)
    all_researcher_ids = [rid for rid in all_researcher_ids if rid]
    
    stats_researchers = Researcher.objects.filter(id__in=all_researcher_ids).distinct().count()

    top_researchers_qs = Researcher.objects.filter(
        id__in=all_researcher_ids
    ).annotate(
        relevant_pubs=Count('articleauthor', filter=models.Q(articleauthor__article__country__in=countries_qs)) + 
                      Count('bookauthor', filter=models.Q(bookauthor__book__country__in=countries_qs))
    ).order_by('-relevant_pubs')[:5]
    
    top_researchers = [
        {
            "name": r.name.title(),
            "category": r.category or ""
        }
        for r in top_researchers_qs
    ]

    grp_ids_art = Article.objects.filter(country__in=countries_qs).values_list('group_id', flat=True)
    
    grp_ids_bk = Book.objects.filter(country__in=countries_qs).values_list('group_id', flat=True)
    
    grp_ids_evt = ScientificEvent.objects.filter(country_obj__in=countries_qs).values_list('group_id', flat=True)
    
    all_grp_ids = set(list(grp_ids_art) + list(grp_ids_bk) + list(grp_ids_evt))
    
    research_lines = ResearchLine.objects.filter(
        groups__id__in=all_grp_ids
    ).values_list('name', flat=True).distinct().order_by('name')

    areas = list(research_lines)
    
    if not areas:
        top_groups_qs = ResearchGroup.objects.filter(id__in=all_grp_ids).annotate(
            count=Count('articles', filter=models.Q(articles__country__in=countries_qs)) + 
                  Count('books', filter=models.Q(books__country__in=countries_qs))
        ).order_by('-count')[:3]
        areas = [g.name for g in top_groups_qs]

    def clean_text(t):
        if not t: return ""
        t = t.replace('\x93', '"').replace('\x94', '"')
        t = t.replace('', '"').replace('', '"')
        t = t.replace('“', '"').replace('”', '"')
        return t

    _r_articles = []
    for a in recent_articles:
        if 'title' in a: a['title'] = clean_text(a['title'])
        _r_articles.append(a)

    _r_books = []
    for b in recent_books:
        if 'title' in b: b['title'] = clean_text(b['title'])
        _r_books.append(b)

    return Response({
        "name": display_name or get_country_display_name(country_id),
        "info": "",
        "stats": {
            "articles": articles_count,
            "books": books_count,
            "chapters": chapters_count,
            "thesis": thesis_count,
            "researchers": stats_researchers
        },
        "topResearchers": top_researchers,
        "areas": areas,
        "chart": [
            { "name": 'Articulos', "value": articles_count, "color": '#10b981' },
            { "name": 'Libros', "value": books_count, "color": '#3b82f6' },
            { "name": 'Capítulos', "value": chapters_count, "color": '#f59e0b' },
        ],
        "recentArticles": _r_articles,
        "recentBooks": _r_books
    })



